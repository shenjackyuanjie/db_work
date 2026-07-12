import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

doc = Document()

# Style setup
style = doc.styles['Normal']
font = style.font
font.name = '宋体'
font.size = Pt(11)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

# Read markdown
with open('柑橘果园智能诊断与管理系统设计与实现_报告.md', 'r', encoding='utf-8') as f:
    lines = f.readlines()

i = 0
in_code_block = False
in_table = False
code_lines = []
code_lang = ''
table_rows = []

def add_paragraph_with_bold(doc, text, font_size=11, bold=False,
                            alignment=None, space_after=Pt(6)):
    p = doc.add_paragraph()
    if alignment is not None:
        p.alignment = alignment
    p.paragraph_format.space_after = space_after
    if not text:
        return p
    # Parse inline bold **text**
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            run.bold = True
            run.font.size = Pt(font_size)
            run.font.name = '宋体'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        else:
            run = p.add_run(part)
            run.font.size = Pt(font_size)
            run.font.name = '宋体'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    return p


def flush_table(doc, rows):
    if not rows:
        return
    # Remove separator lines like |---|---|
    clean_rows = [r for r in rows if not re.match(r'^\|[\s\-:|]+\|$', r)]
    if len(clean_rows) < 1:
        return
    cols = len([c for c in clean_rows[0].split('|') if c.strip() != ''])
    table = doc.add_table(rows=len(clean_rows), cols=cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(clean_rows):
        cells = [c.strip() for c in row.split('|') if c.strip() != '']
        for ci, cell_text in enumerate(cells):
            if ci < cols:
                cell = table.rows[ri].cells[ci]
                cell.text = ''
                p = cell.paragraphs[0]
                run = p.add_run(cell_text)
                run.font.size = Pt(9)
                run.font.name = '宋体'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
                if ri == 0:
                    run.bold = True
    doc.add_paragraph()  # spacer after table


def flush_code_block(doc, lines, lang):
    if not lines:
        return
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Cm(1)
        run = p.add_run(line)
        run.font.name = 'Consolas'
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    doc.add_paragraph()


while i < len(lines):
    line = lines[i].rstrip()

    # Code block toggle
    if line.startswith('```'):
        if in_code_block:
            flush_code_block(doc, code_lines, code_lang)
            code_lines = []
            code_lang = ''
            in_code_block = False
        else:
            in_code_block = True
            code_lang = line[3:].strip()
        i += 1
        continue

    if in_code_block:
        code_lines.append(line)
        i += 1
        continue

    # Table handling
    if line.startswith('|') and line.endswith('|'):
        if not in_table:
            in_table = True
            table_rows = []
        table_rows.append(line)
        i += 1
        continue
    else:
        if in_table:
            flush_table(doc, table_rows)
            table_rows = []
            in_table = False

    # Heading
    if line.startswith('# '):
        add_paragraph_with_bold(doc, line[2:], font_size=22, bold=True,
                                alignment=WD_ALIGN_PARAGRAPH.CENTER,
                                space_after=Pt(12))
    elif line.startswith('## '):
        add_paragraph_with_bold(doc, line[3:], font_size=16, bold=True,
                                space_after=Pt(8))
    elif line.startswith('### '):
        add_paragraph_with_bold(doc, line[4:], font_size=13, bold=True,
                                space_after=Pt(6))
    elif line.startswith('- '):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1)
        p.paragraph_format.space_after = Pt(3)
        text = line[2:]
        parts = re.split(r'(\*\*.*?\*\*)', text)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part[2:-2])
                run.bold = True
                run.font.size = Pt(11)
                run.font.name = '宋体'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
            else:
                run = p.add_run(part)
                run.font.size = Pt(11)
                run.font.name = '宋体'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    elif line.strip() == '':
        pass
    else:
        add_paragraph_with_bold(doc, line, font_size=11)

    i += 1

# Flush remaining
if in_table:
    flush_table(doc, table_rows)
if in_code_block:
    flush_code_block(doc, code_lines, code_lang)

# Set page margins
for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

output_path = '柑橘果园智能诊断与管理系统设计与实现_报告.docx'
doc.save(output_path)
print(f'Done! Saved to: {output_path}')
